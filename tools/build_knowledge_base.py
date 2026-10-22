from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "reference" / "【可任意分享】健身Excel超级套表（作者：B站好人松松）26年9月最新版.xlsx"
DEFAULT_OUTPUT = ROOT / "knowledge-base"


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def slugify(name: str) -> str:
    text = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "-", name).strip("-").lower()
    return text[:70] or "sheet"


def source_ref(sheet: str, address: str) -> dict[str, str]:
    return {"sheet": sheet, "cell": address}


def color_value(color: Any) -> str | None:
    if color is None:
        return None
    if getattr(color, "type", None) == "rgb":
        return color.rgb
    if getattr(color, "type", None) == "indexed":
        return f"indexed:{color.indexed}"
    if getattr(color, "type", None) == "theme":
        return f"theme:{color.theme}:{getattr(color, 'tint', 0)}"
    return None


def anchor_info(anchor: Any) -> dict[str, Any]:
    if hasattr(anchor, "_from"):
        result = {
            "from": {
                "row": anchor._from.row + 1,
                "column": anchor._from.col + 1,
                "cell": f"{get_column_letter(anchor._from.col + 1)}{anchor._from.row + 1}",
            }
        }
        if hasattr(anchor, "to"):
            result["to"] = {
                "row": anchor.to.row + 1,
                "column": anchor.to.col + 1,
                "cell": f"{get_column_letter(anchor.to.col + 1)}{anchor.to.row + 1}",
            }
        return result
    return {"value": str(anchor)}


def cell_record(cell: Any, cached_cell: Any) -> dict[str, Any]:
    value = json_value(cell.value)
    is_formula = isinstance(value, str) and value.startswith("=")
    result: dict[str, Any] = {
        "address": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "value": None if is_formula else value,
        "formula": value if is_formula else None,
        "cached_value": json_value(cached_cell.value) if is_formula else None,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "style_id": cell.style_id,
    }
    if cell.hyperlink:
        result["hyperlink"] = {
            "target": cell.hyperlink.target,
            "location": cell.hyperlink.location,
            "display": cell.hyperlink.display,
        }
    if cell.comment:
        result["comment"] = {"author": cell.comment.author, "text": cell.comment.text}
    if cell.has_style:
        result["style"] = {
            "font": {
                "name": cell.font.name,
                "size": cell.font.sz,
                "bold": cell.font.b,
                "italic": cell.font.i,
                "color": color_value(cell.font.color),
            },
            "fill": color_value(cell.fill.fgColor),
            "alignment": {
                "horizontal": cell.alignment.horizontal,
                "vertical": cell.alignment.vertical,
                "wrap_text": cell.alignment.wrap_text,
                "text_rotation": cell.alignment.text_rotation,
            },
        }
    return result


def row_values(ws: Any, row: int, min_col: int = 1, max_col: int | None = None) -> list[dict[str, Any]]:
    max_col = max_col or ws.max_column
    values = []
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row, col)
        if cell.value not in (None, ""):
            values.append({"address": cell.coordinate, "value": json_value(cell.value)})
    return values


def extract_foods(ws: Any) -> dict[str, Any]:
    carbs: list[dict[str, Any]] = []
    protein: list[dict[str, Any]] = []
    guidance: list[dict[str, Any]] = []
    category = None
    for row in range(5, 92):
        if ws.cell(row, 2).value:
            category = str(ws.cell(row, 2).value).strip()
        food = ws.cell(row, 3).value
        if not food:
            continue
        rate = ws.cell(row, 4).value
        carbs.append({
            "category": category,
            "food": str(food).strip(),
            "carbohydrate_rate": json_value(rate),
            "gi": json_value(ws.cell(row, 5).value),
            "notes": json_value(ws.cell(row, 6).value),
            "source": source_ref(ws.title, f"B{row}:F{row}"),
        })
    category = None
    for row in range(96, 117):
        if ws.cell(row, 2).value:
            category = str(ws.cell(row, 2).value).strip()
        food = ws.cell(row, 3).value
        if not food:
            continue
        protein.append({
            "category": category,
            "food": str(food).strip(),
            "protein_rate": json_value(ws.cell(row, 4).value),
            "positioning": json_value(ws.cell(row, 5).value),
            "notes": json_value(ws.cell(row, 6).value),
            "source": source_ref(ws.title, f"B{row}:F{row}"),
        })
    for row in range(118, ws.max_row + 1):
        values = row_values(ws, row, 2, 6)
        if values:
            guidance.append({"row": row, "cells": values, "source": source_ref(ws.title, f"B{row}:F{row}")})
    return {
        "schema_version": "1.0.0",
        "source_sheet": ws.title,
        "carbohydrate_foods": carbs,
        "protein_foods": protein,
        "fat_and_selection_guidance": guidance,
    }


APP_FOODS = [
    {"name": "馒头 / 花卷", "kcal": 223, "carb": 47, "protein": 7, "fat": 1, "unit": "g", "reference_note": "参考表碳水率约 50%", "source_cell": "D15"},
    {"name": "米饭（一般）", "kcal": 116, "carb": 26, "protein": 2.6, "fat": 0.3, "unit": "g", "reference_note": "参考表碳水率约 30%", "source_cell": "D7"},
    {"name": "米饭（偏硬）", "kcal": 135, "carb": 30, "protein": 2.8, "fat": 0.3, "unit": "g", "reference_note": "参考表碳水率约 35%", "source_cell": "D8"},
    {"name": "红薯（蒸煮）", "kcal": 86, "carb": 20, "protein": 1.6, "fat": 0.1, "unit": "g", "reference_note": "参考表碳水率约 18%", "source_cell": "D24"},
    {"name": "土豆（蒸煮）", "kcal": 77, "carb": 17, "protein": 2, "fat": 0.1, "unit": "g", "reference_note": "参考表碳水率约 18%", "source_cell": "D26"},
    {"name": "全麦切片面包", "kcal": 247, "carb": 41, "protein": 13, "fat": 4.2, "unit": "g", "reference_note": "参考表切片面包碳水率约 50%，品牌营养表优先", "source_cell": "D14"},
    {"name": "燕麦片（生）", "kcal": 379, "carb": 60, "protein": 13.2, "fat": 6.5, "unit": "g", "reference_note": "参考表碳水率约 60%", "source_cell": "D36"},
    {"name": "熟面条", "kcal": 138, "carb": 25, "protein": 4.5, "fat": 2, "unit": "g", "reference_note": "参考表细面碳水率约 23%", "source_cell": "D22"},
    {"name": "鸡胸肉（熟）", "kcal": 165, "carb": 0, "protein": 31, "fat": 3.6, "unit": "g", "reference_note": "参考表熟瘦肉蛋白率约 25–30%", "source_cell": "D99"},
    {"name": "鱼虾（生）", "kcal": 100, "carb": 0, "protein": 18, "fat": 2, "unit": "g", "reference_note": "参考表鱼虾蛋白率约 15%", "source_cell": "D96"},
    {"name": "瘦牛肉（熟）", "kcal": 200, "carb": 0, "protein": 27, "fat": 9, "unit": "g", "reference_note": "参考表一般熟瘦肉蛋白率约 25%", "source_cell": "D99"},
    {"name": "鸡蛋", "kcal": 72, "carb": 0.4, "protein": 6, "fat": 5, "unit": "serving", "grams": 50, "reference_note": "每个蛋白质 6g、脂肪 5g", "source_cell": "D109"},
    {"name": "纯牛奶 250ml", "kcal": 150, "carb": 12, "protein": 10, "fat": 9, "unit": "serving", "grams": 250, "reference_note": "每盒 250ml", "source_cell": "D110"},
    {"name": "无糖酸奶 200ml", "kcal": 120, "carb": 10, "protein": 10, "fat": 5, "unit": "serving", "grams": 200, "reference_note": "以包装营养表为准", "source_cell": "D111"},
    {"name": "豆腐", "kcal": 81, "carb": 3, "protein": 8, "fat": 4.5, "unit": "g", "reference_note": "参考表蛋白率约 7%", "source_cell": "D115"},
    {"name": "毛豆", "kcal": 131, "carb": 11, "protein": 13, "fat": 5, "unit": "g", "reference_note": "参考表蛋白率约 13%", "source_cell": "D113"},
    {"name": "苹果", "kcal": 52, "carb": 14, "protein": 0.3, "fat": 0.2, "unit": "g", "reference_note": "水果应计入碳水", "source_cell": "C50"},
    {"name": "香蕉", "kcal": 89, "carb": 23, "protein": 1.1, "fat": 0.3, "unit": "g", "reference_note": "水果应计入碳水", "source_cell": "C50"},
    {"name": "西兰花", "kcal": 34, "carb": 6.6, "protein": 2.8, "fat": 0.4, "unit": "g", "reference_note": "蔬菜无需过度精算", "source_cell": "B120"},
    {"name": "坚果", "kcal": 607, "carb": 21, "protein": 20, "fat": 54, "unit": "g", "reference_note": "高脂食物，建议称重", "source_cell": "C120"},
    {"name": "乳清蛋白粉", "kcal": 390, "carb": 8, "protein": 75, "fat": 6, "unit": "g", "reference_note": "参考表蛋白率约 70–80%", "source_cell": "D102"},
]


def semantic_rows(ws: Any, start: int = 1, end: int | None = None) -> list[dict[str, Any]]:
    records = []
    for row in range(start, (end or ws.max_row) + 1):
        values = row_values(ws, row)
        if values:
            records.append({
                "row": row,
                "cells": values,
                "text": " | ".join(f"{x['address']} {x['value']}" for x in values),
                "source": source_ref(ws.title, f"A{row}:{get_column_letter(ws.max_column)}{row}"),
            })
    return records


def extract_faq(ws: Any) -> dict[str, Any]:
    entries = []
    for row in range(1, ws.max_row + 1):
        vals = [ws.cell(row, c).value for c in range(1, min(ws.max_column, 3) + 1)]
        text_vals = [str(v).strip() for v in vals if v not in (None, "")]
        if not text_vals:
            continue
        question = next((v for v in text_vals if "？" in v or "?" in v or re.match(r"^\d+[\.、]", v)), text_vals[0])
        answer_candidates = [v for v in text_vals if v != question]
        entries.append({
            "row": row,
            "question_or_heading": question,
            "answer": "\n".join(answer_candidates) or None,
            "cells": row_values(ws, row, 1, min(ws.max_column, 3)),
            "source": source_ref(ws.title, f"A{row}:C{row}"),
        })
    return {"source_sheet": ws.title, "entries": entries}


def extract_one_rm(ws: Any) -> dict[str, Any]:
    formulas = []
    for row in range(7, ws.max_row + 1):
        if ws.cell(row, 2).value:
            formulas.append({
                "author": json_value(ws.cell(row, 2).value),
                "description": json_value(ws.cell(row, 3).value),
                "input_weight_example_kg": json_value(ws.cell(row, 4).value),
                "input_repetitions_example": json_value(ws.cell(row, 5).value),
                "excel_formula": json_value(ws.cell(row, 6).value),
                "accuracy_note": json_value(ws.cell(row, 7).value),
                "source": source_ref(ws.title, f"B{row}:G{row}"),
            })
    return {
        "source_sheet": ws.title,
        "title": json_value(ws["B2"].value),
        "instructions": json_value(ws["B4"].value),
        "literature_source": json_value(ws["B5"].value),
        "formulas": formulas,
    }


def build(source: Path, output: Path) -> None:
    source = source.resolve()
    output = output.resolve()
    if output.exists():
        raise FileExistsError(f"Output already exists: {output}. Remove it explicitly before rebuilding.")
    if ROOT not in output.parents:
        raise ValueError("Output must stay inside the workspace.")

    raw_dir = output / "source" / "sheets"
    entity_dir = output / "entities"
    search_dir = output / "search"
    media_dir = output / "media"
    raw_media_dir = media_dir / "raw"
    for path in (raw_dir, entity_dir, search_dir, raw_media_dir):
        path.mkdir(parents=True, exist_ok=True)

    raw_hash_to_file: dict[str, str] = {}
    with zipfile.ZipFile(source) as archive:
        for member in archive.infolist():
            if member.filename.startswith("xl/media/") and not member.is_dir():
                data = archive.read(member)
                target = raw_media_dir / Path(member.filename).name
                target.write_bytes(data)
                raw_hash_to_file[hashlib.sha256(data).hexdigest()] = f"media/raw/{target.name}"

    wb = openpyxl.load_workbook(source, read_only=False, data_only=False)
    cached = openpyxl.load_workbook(source, read_only=False, data_only=True)
    sheet_catalog = []
    chunks = []
    total_cells = 0
    image_catalog = []

    for index, ws in enumerate(wb.worksheets):
        cached_ws = cached[ws.title]
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None or cell.hyperlink or cell.comment:
                    cells.append(cell_record(cell, cached_ws[cell.coordinate]))
        total_cells += len(cells)

        rows = semantic_rows(ws)
        for row in rows:
            chunks.append({
                "id": f"sheet-{index:02d}-row-{row['row']:04d}",
                "sheet_index": index,
                "sheet": ws.title,
                "range": row["source"]["cell"],
                "content": row["text"],
                "kind": "worksheet_row",
            })

        images = []
        for image_index, image in enumerate(getattr(ws, "_images", []), start=1):
            extension = (getattr(image, "format", None) or "png").lower()
            filename = f"sheet-{index:02d}-image-{image_index:03d}.{extension}"
            target = media_dir / filename
            try:
                data = image._data()
                digest = hashlib.sha256(data).hexdigest()
                mapped_file = raw_hash_to_file.get(digest)
                if mapped_file:
                    file_path = mapped_file
                else:
                    target.write_bytes(data)
                    file_path = f"media/{filename}"
                record = {
                    "file": file_path,
                    "width_px": getattr(image, "width", None),
                    "height_px": getattr(image, "height", None),
                    "anchor": anchor_info(image.anchor),
                    "sha256": digest,
                }
            except Exception as exc:
                record = {"error": str(exc), "anchor": anchor_info(image.anchor)}
            images.append(record)
            image_catalog.append({"sheet": ws.title, **record})

        validations = []
        for validation in ws.data_validations.dataValidation:
            validations.append({
                "ranges": str(validation.sqref),
                "type": validation.type,
                "operator": validation.operator,
                "formula1": json_value(validation.formula1),
                "formula2": json_value(validation.formula2),
                "allow_blank": validation.allow_blank,
            })

        conditional_formats = []
        for key in ws.conditional_formatting:
            conditional_formats.append({
                "range": str(key.sqref),
                "rules": [{"type": rule.type, "operator": getattr(rule, "operator", None), "formula": list(rule.formula or [])} for rule in ws.conditional_formatting[key]],
            })

        raw = {
            "schema_version": "1.0.0",
            "sheet_index": index,
            "title": ws.title,
            "state": ws.sheet_state,
            "dimensions": {"max_row": ws.max_row, "max_column": ws.max_column},
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
            "row_dimensions": {str(k): {"height": v.height, "hidden": v.hidden} for k, v in ws.row_dimensions.items() if v.height or v.hidden},
            "column_dimensions": {k: {"width": v.width, "hidden": v.hidden} for k, v in ws.column_dimensions.items() if v.width or v.hidden},
            "data_validations": validations,
            "conditional_formats": conditional_formats,
            "images": images,
            "charts": [{"type": type(chart).__name__, "title": str(getattr(chart, "title", "")), "anchor": anchor_info(chart.anchor)} for chart in getattr(ws, "_charts", [])],
            "cells": cells,
        }
        filename = f"{index:02d}-{slugify(ws.title)}.json"
        (raw_dir / filename).write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
        sheet_catalog.append({
            "index": index,
            "title": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "nonempty_cells": len(cells),
            "images": len(images),
            "charts": len(getattr(ws, "_charts", [])),
            "raw_file": f"source/sheets/{filename}",
        })

    food_data = extract_foods(wb["19日常食物营养率"])
    for item in APP_FOODS:
        item["source"] = source_ref("19日常食物营养率", item.pop("source_cell"))
        item["nutrient_values_note"] = "碳水或蛋白比例参考原表；总热量及原表未提供的营养素为产品层估算，包装营养表优先。"
    food_data["application_foods"] = APP_FOODS
    (entity_dir / "foods.json").write_text(json.dumps(food_data, ensure_ascii=False, indent=2), encoding="utf-8")

    meal_plans = []
    for ws in wb.worksheets[1:16]:
        meal_plans.append({
            "sheet_index": wb.sheetnames.index(ws.title),
            "name": ws.title,
            "goal": "fat_loss" if "减脂" in ws.title else "muscle_gain",
            "training_time": next((x for x in ["早饭后练", "午饭前练", "午饭后练", "晚饭前练", "晚饭后练", "夜里练", "无力训"] if x in ws.title), None),
            "profile_and_rules": semantic_rows(ws, 4, 25),
            "meal_structure": semantic_rows(ws, 27, ws.max_row),
            "raw_sheet_file": sheet_catalog[wb.sheetnames.index(ws.title)]["raw_file"],
        })
    (entity_dir / "meal-plans.json").write_text(json.dumps({"plans": meal_plans}, ensure_ascii=False, indent=2), encoding="utf-8")

    aerobic_ws = wb["16有氧热量消耗"]
    (entity_dir / "aerobic-energy.json").write_text(json.dumps({"source_sheet": aerobic_ws.title, "rows": semantic_rows(aerobic_ws)}, ensure_ascii=False, indent=2), encoding="utf-8")
    faq_data = {"fat_loss": extract_faq(wb["17减脂-问答汇总"]), "muscle_gain": extract_faq(wb["18增肌-问答汇总"])}
    (entity_dir / "faqs.json").write_text(json.dumps(faq_data, ensure_ascii=False, indent=2), encoding="utf-8")

    products_ws = wb["20蛋白质价格比较"]
    (entity_dir / "protein-products.json").write_text(json.dumps({"source_sheet": products_ws.title, "rows": semantic_rows(products_ws)}, ensure_ascii=False, indent=2), encoding="utf-8")
    training = []
    for ws in wb.worksheets[21:25]:
        training.append({"name": ws.title, "rows": semantic_rows(ws), "raw_sheet_file": sheet_catalog[wb.sheetnames.index(ws.title)]["raw_file"]})
    (entity_dir / "training-plans.json").write_text(json.dumps({"plans": training}, ensure_ascii=False, indent=2), encoding="utf-8")
    one_rm = extract_one_rm(wb["25最大力量预测公式"])
    (entity_dir / "one-rep-max.json").write_text(json.dumps(one_rm, ensure_ascii=False, indent=2), encoding="utf-8")

    anatomy = []
    for ws in wb.worksheets[26:31]:
        anatomy.append({
            "name": ws.title,
            "rows": semantic_rows(ws),
            "images": [image for image in image_catalog if image["sheet"] == ws.title],
            "raw_sheet_file": sheet_catalog[wb.sheetnames.index(ws.title)]["raw_file"],
        })
    (entity_dir / "anatomy-and-stretching.json").write_text(json.dumps({"sections": anatomy}, ensure_ascii=False, indent=2), encoding="utf-8")

    rules = {
        "schema_version": "1.0.0",
        "rules": [
            {"id": "bmr-mifflin-st-jeor", "topic": "energy", "description": "男性基础代谢=体重kg×9.99+身高cm×6.25-年龄×4.92+5；女性将常数项改为-161。", "source": source_ref("1减脂-早饭后练（早起版）", "G13")},
            {"id": "non-exercise-energy", "topic": "energy", "description": "原表以基础代谢约占无运动总消耗70%估算无运动总消耗。", "source": source_ref("1减脂-早饭后练（早起版）", "D14:G14")},
            {"id": "fat-loss-rate", "topic": "weight", "description": "原表建议以至少两周趋势评估减重，单日体重容易受食物与水分影响。", "source": source_ref("1减脂-早饭后练（早起版）", "B9:C9")},
            {"id": "strength-frequency-fat-loss", "topic": "training", "description": "减脂期为减少肌肉损失，原表建议每周力量训练3–5次。", "source": source_ref("1减脂-早饭后练（早起版）", "C24:E24")},
            {"id": "strength-frequency-gain", "topic": "training", "description": "增肌需要稳定力量训练，原表建议每周3–5次。", "source": source_ref("9增肌-早饭后练（早起版）", "C24:E24")},
            {"id": "aerobic-weekly-average", "topic": "training", "description": "饮食方案中的每日有氧消耗=每小时消耗×每周小时数÷7。", "source": source_ref("16有氧热量消耗", "B6:F7")},
            {"id": "fruit-counts-as-carbs", "topic": "nutrition", "description": "水果中的糖属于碳水，应计入碳水配额。", "source": source_ref("1减脂-早饭后练（早起版）", "N46:N53")},
            {"id": "vegetables-not-overtracked", "topic": "nutrition", "description": "原表认为普通蔬菜碳水率较低，不必过度定量。", "source": source_ref("1减脂-早饭后练（早起版）", "N31:N34")},
        ],
        "medical_disclaimer": "该知识库来自作者提供的健身经验表格，仅用于产品参考与自我记录，不构成医疗诊断、治疗或个体化营养处方。",
    }
    (entity_dir / "rules.json").write_text(json.dumps(rules, ensure_ascii=False, indent=2), encoding="utf-8")

    with (search_dir / "chunks.jsonl").open("w", encoding="utf-8") as handle:
        for chunk in chunks:
            handle.write(json.dumps(chunk, ensure_ascii=False) + "\n")

    entity_index = [
        {"id": "foods", "file": "entities/foods.json", "description": "碳水、蛋白质、脂肪指导及应用食物库"},
        {"id": "meal_plans", "file": "entities/meal-plans.json", "description": "15套减脂/增肌饮食方案"},
        {"id": "aerobic_energy", "file": "entities/aerobic-energy.json", "description": "有氧热量消耗方法与查表数据"},
        {"id": "faqs", "file": "entities/faqs.json", "description": "减脂与增肌问答"},
        {"id": "protein_products", "file": "entities/protein-products.json", "description": "蛋白质产品与价格比较"},
        {"id": "training_plans", "file": "entities/training-plans.json", "description": "健身房与居家训练计划"},
        {"id": "one_rep_max", "file": "entities/one-rep-max.json", "description": "1RM预测公式"},
        {"id": "anatomy", "file": "entities/anatomy-and-stretching.json", "description": "拉伸、肌肉与关节活动资料"},
        {"id": "rules", "file": "entities/rules.json", "description": "跨模块可执行规则及来源"},
    ]
    catalog = {"schema_version": "1.0.0", "entities": entity_index, "sheets": sheet_catalog, "images": image_catalog}
    (output / "catalog.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")

    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    manifest = {
        "name": "Boji Health Fitness Knowledge Base",
        "schema_version": "1.0.0",
        "language": "zh-CN",
        "generated_at": datetime.now().astimezone().isoformat(),
        "source": {"file": source.name, "sha256": source_hash},
        "coverage": {
            "worksheets": len(wb.worksheets),
            "nonempty_cells": total_cells,
            "search_chunks": len(chunks),
            "mapped_images": len(image_catalog),
            "raw_media_files": len(list(raw_media_dir.iterdir())),
            "semantic_entities": len(entity_index),
        },
        "entrypoints": {"catalog": "catalog.json", "browser_bundle": "browser-bundle.js", "search_chunks": "search/chunks.jsonl"},
    }
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    browser_bundle = {"manifest": manifest, "foods": APP_FOODS, "rules": rules["rules"], "catalog": entity_index}
    (output / "browser-bundle.js").write_text("window.BOJI_KB = " + json.dumps(browser_bundle, ensure_ascii=False) + ";\n", encoding="utf-8")

    readme = f"""# Boji Health 知识库

本目录由 `tools/build_knowledge_base.py` 从参考 Excel 自动生成。原始工作簿未被修改。

## 结构

- `manifest.json`：版本、来源校验值与覆盖统计。
- `catalog.json`：31 个工作表、语义实体和图片索引。
- `source/sheets/`：每个工作表的完整结构化快照，保留非空单元格、公式、缓存值、样式、合并区域、验证规则、图片锚点及来源坐标。
- `entities/`：供产品直接调用的食品、饮食方案、有氧、问答、训练计划、1RM、拉伸解剖和跨模块规则。
- `search/chunks.jsonl`：按工作表行生成的检索块，可直接导入全文检索或向量数据库。
- `media/`：按工作表映射的图片；`media/raw/` 保留工作簿压缩包中的全部原始媒体。
- `browser-bundle.js`：无需服务器即可由前端 `<script>` 引用的精简知识库。

## 来源与边界

来源：`{source.name}`
SHA-256：`{source_hash}`

知识库保留原作者表述和公式，并通过 `source.sheet` / `source.cell` 标注来源。食品总热量与原表未提供的营养素在 `application_foods` 中明确标为产品层估算，实际使用时应优先采用食品包装营养表。该资料不构成医疗建议。
"""
    (output / "README.md").write_text(readme, encoding="utf-8")

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
